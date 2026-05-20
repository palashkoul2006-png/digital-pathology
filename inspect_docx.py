"""
inspect_docx.py - Quick audit of Digital_Pathology_Diagrams.docx
Prints headings, image count, and paragraph summary.
"""
import os, glob
from docx import Document

# ── Search for the file in likely locations ───────────────────────────────────
SEARCH_DIRS = [
    r"C:\Users\palas\OneDrive - Vishwakarma Institute of Technology\Desktop\digital_pathology",
    r"C:\Users\palas\Desktop",
    r"C:\Users\palas\Documents",
    r"C:\Users\palas",
]

DOCX_PATH = None
for d in SEARCH_DIRS:
    candidate = os.path.join(d, "Digital_Pathology_Diagrams.docx")
    if os.path.exists(candidate):
        DOCX_PATH = candidate
        print(f"✅ Found file at: {DOCX_PATH}")
        print(f"   File size    : {os.path.getsize(DOCX_PATH) / 1024:.1f} KB\n")
        break

# Broader glob search if not found in expected dirs
if not DOCX_PATH:
    print("⚠️  Not found in expected locations. Running broad search...")
    matches = glob.glob(r"C:\Users\palas\**\Digital_Pathology_Diagrams.docx", recursive=True)
    if matches:
        DOCX_PATH = matches[0]
        print(f"✅ Found at: {DOCX_PATH}\n")
    else:
        print("❌ File not found anywhere under C:\\Users\\palas\\")
        print("   Please re-run: python export_diagrams.py")
        exit(1)

doc = Document(DOCX_PATH)

# ── Count inline images ───────────────────────────────────────────────────────
image_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_count += 1

# ── Collect headings and captions ────────────────────────────────────────────
headings  = []
captions  = []
body_paras = 0

for para in doc.paragraphs:
    style = para.style.name
    text  = para.text.strip()
    if not text:
        continue
    if style.startswith('Heading'):
        level = style.replace('Heading ', 'H')
        headings.append((level, text))
    elif 'Figure:' in text or text.startswith('Figure'):
        captions.append(text)
    else:
        body_paras += 1

# ── Report ────────────────────────────────────────────────────────────────────
print("=" * 60)
print("  DOCX INSPECTION REPORT")
print("=" * 60)
print(f"\n📊 Summary")
print(f"   Total paragraphs : {len(doc.paragraphs)}")
print(f"   Headings found   : {len(headings)}")
print(f"   Images embedded  : {image_count}")
print(f"   Diagram captions : {len(captions)}")

print(f"\n📑 Document Structure (Headings)")
print("-" * 50)
for level, text in headings:
    indent = "  " * (int(level[1]) - 1)
    print(f"   {indent}[{level}] {text}")

print(f"\n🖼️  Diagram Captions")
print("-" * 50)
for i, cap in enumerate(captions, 1):
    status = "✅" if len(cap) > 10 else "⚠️"
    print(f"   {status} [{i:02d}] {cap[:80]}")

print(f"\n{'='*60}")
missing = 11 - image_count
if missing <= 0:
    print("✅ All 11 diagrams are present in the document!")
else:
    print(f"⚠️  {missing} diagram(s) missing (expected 11, found {image_count})")
print("=" * 60)
