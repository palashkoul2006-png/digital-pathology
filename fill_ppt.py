"""
fill_ppt.py — Fill the VIT Pune National Competition PPT template
with PathAI Digital Pathology project content.
Preserves all template design elements (groups, shapes, colors).
"""
import os, copy
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(PROJECT_DIR, "PPT TEMPLATE - National Level Project Competition - VIT, Pune (1).pptx")
OUTPUT = os.path.join(PROJECT_DIR, "PathAI_Presentation.pptx")

# ── Font helpers ─────────────────────────────────────────────────────────────
HEADING_SIZE = Emu(685546)   # ~54pt — matches template headings
BODY_BOLD    = Emu(533400)   # ~42pt — matches template body bold
BODY_NORM    = Emu(533400)   # ~42pt
SMALL_BODY   = Emu(406146)   # ~32pt
DETAIL_SIZE  = Emu(355600)   # ~28pt
BULLET_SIZE  = Emu(330000)   # ~26pt
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_RED = RGBColor(0x8B, 0x00, 0x00)
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)

def clear_text(shape):
    """Remove all text from a shape's text frame."""
    tf = shape.text_frame
    for para in tf.paragraphs:
        for run in para.runs:
            run.text = ""
    # Remove all but first paragraph
    while len(tf.paragraphs) > 1:
        p_elem = tf.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)

def add_line(tf, text, size=BODY_NORM, bold=False, color=BLACK, first=False):
    """Add a formatted paragraph to a text frame."""
    if first and tf.paragraphs and tf.paragraphs[0].text == "":
        para = tf.paragraphs[0]
    else:
        para = tf.add_paragraph()
    run = para.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return para

def add_heading_line(tf, text, first=False):
    return add_line(tf, text, size=BODY_BOLD, bold=True, color=BLACK, first=first)

def add_bullet(tf, text, indent=0):
    p = add_line(tf, f"• {text}", size=BULLET_SIZE, bold=False, color=BLACK)
    return p

def add_sub_bullet(tf, text):
    return add_line(tf, f"   ‣ {text}", size=Emu(300000), bold=False, color=BLACK)

def find_textbox(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None

def add_textbox(slide, left, top, width, height):
    from pptx.util import Emu as E
    return slide.shapes.add_textbox(E(left), E(top), E(width), E(height))

# ── Load template ────────────────────────────────────────────────────────────
prs = Presentation(TEMPLATE)
slides = list(prs.slides)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE PAGE
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 1: Title Page")
s1 = slides[0]

tb25 = find_textbox(s1, 'TextBox 25')
if tb25:
    clear_text(tb25)
    add_line(tb25.text_frame, "PathAI", size=HEADING_SIZE, bold=True, color=BLACK, first=True)

tb26 = find_textbox(s1, 'TextBox 26')
if tb26:
    clear_text(tb26)
    tf = tb26.text_frame
    add_line(tf, "", size=BODY_BOLD, first=True)
    add_line(tf, "Problem Statement Title — AI-Powered Skin Cancer Detection", size=BODY_BOLD, bold=True)
    add_line(tf, "    & Clinical Decision Support System", size=BODY_BOLD, bold=True)
    add_line(tf, "", size=BODY_BOLD)
    add_line(tf, "Domain — Healthcare / Artificial Intelligence / Deep Learning", size=BODY_BOLD, bold=True)
    add_line(tf, "", size=BODY_BOLD)
    add_line(tf, "Team Name — [YOUR TEAM NAME]", size=BODY_BOLD, bold=True)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — EXISTING SYSTEM AND DRAWBACKS
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 2: Existing System & Drawbacks")
s2 = slides[1]

content_box = add_textbox(s2, 1790710, 2200000, 14706581, 7000000)
tf = content_box.text_frame
tf.word_wrap = True

add_heading_line(tf, "Current Diagnostic Process", first=True)
add_bullet(tf, "Manual visual inspection by dermatologists — subjective, prone to human error")
add_bullet(tf, "Dermoscopy requires trained specialists — limited availability in rural areas")
add_bullet(tf, "Biopsy-based confirmation takes days/weeks — delays critical treatment")
add_bullet(tf, "Misdiagnosis rate: 25–30% for early-stage melanoma even by experts")
add_line(tf, "", size=BULLET_SIZE)

add_heading_line(tf, "Drawbacks of Existing AI Solutions")
add_bullet(tf, "Most use outdated CNNs (ResNet/VGG) with lower accuracy (<90%)")
add_bullet(tf, "No explainability — black-box predictions reduce doctor trust")
add_bullet(tf, "No attention heatmaps showing WHERE the model is looking")
add_bullet(tf, "No agentic AI for personalized medical recommendations")
add_bullet(tf, "No progression simulation — cannot predict how lesions may develop")
add_bullet(tf, "No patient history tracking for longitudinal risk monitoring")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — PROPOSED SOLUTION
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 3: Proposed Solution")
s3 = slides[2]

tb12 = find_textbox(s3, 'TextBox 12')
if tb12:
    clear_text(tb12)
    add_line(tb12.text_frame, "PathAI — Digital Pathology", size=HEADING_SIZE, bold=True, first=True)

tb13 = find_textbox(s3, 'TextBox 13')
if tb13:
    clear_text(tb13)
    tf = tb13.text_frame
    add_heading_line(tf, "Proposed Solution (Full Implementation)", first=True)
    add_line(tf, "", size=BULLET_SIZE)
    add_heading_line(tf, "Detailed Breakdown:")
    add_bullet(tf, "Fine-tuned Vision Transformer (ViT) achieving 98.52% accuracy on 7-class skin cancer")
    add_bullet(tf, "ViT Attention Rollout heatmaps — real transformer attention, not Grad-CAM approximation")
    add_bullet(tf, "Gemini 2.5 Flash agentic AI for dynamic, personalized medical reports")
    add_bullet(tf, "Stable Diffusion + custom LoRA for melanoma progression simulation")
    add_line(tf, "", size=BULLET_SIZE)
    add_heading_line(tf, "How It Solves the Problem:")
    add_bullet(tf, "Instant, accurate classification — reduces diagnostic delay from days to seconds")
    add_bullet(tf, "Explainable AI — attention heatmaps show exactly which regions drive the diagnosis")
    add_bullet(tf, "Agentic recommendations — adaptive clinical advice, not static lookup tables")
    add_line(tf, "", size=BULLET_SIZE)
    add_heading_line(tf, "Key Innovations:")
    add_bullet(tf, "First to combine ViT + Attention Rollout + Agentic LLM + Generative Simulation")
    add_bullet(tf, "Custom-trained LoRA on melanoma data for clinical-grade progression prediction")
    add_bullet(tf, "H&E stain auto-detection — classifies pathology vs clinical images automatically")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — TECHNICAL APPROACH (Technologies)
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 4: Technical Approach — Technologies")
s4 = slides[3]

tb13_s4 = find_textbox(s4, 'TextBox 13')
if tb13_s4:
    clear_text(tb13_s4)
    tf = tb13_s4.text_frame
    add_heading_line(tf, "Technologies Involved", first=True)
    add_bullet(tf, "Model: google/vit-base-patch16-224 (Vision Transformer) — HuggingFace Transformers")
    add_bullet(tf, "Framework: PyTorch 2.7 with CUDA GPU acceleration (RTX 3050)")
    add_bullet(tf, "Backend: Flask 3.x (Python 3.11) — REST API server")
    add_bullet(tf, "Generative AI: Stable Diffusion v1.5 + Custom LoRA (diffusers library)")
    add_bullet(tf, "Agentic AI: Google Gemini 2.5 Flash — dynamic medical report generation")
    add_bullet(tf, "Database: MongoDB — patient record persistence & history tracking")
    add_bullet(tf, "Frontend: HTML5, CSS3 (glassmorphism dark UI), Vanilla JS, Chart.js")
    add_bullet(tf, "Image Processing: OpenCV — denoising, H&E stain detection, entropy analysis")
    add_line(tf, "", size=BULLET_SIZE)
    add_heading_line(tf, "Implementation Methodology")
    add_bullet(tf, "Dataset: HAM10000 (10,015 dermatoscopic images, 7 diagnostic categories)")
    add_bullet(tf, "Training: 10 epochs, batch=32, lr=2e-5, FP16 mixed precision, warmup=0.1")
    add_bullet(tf, "Architecture: ViT-Base (86M params) with replaced 7-class classification head")
    add_bullet(tf, "Heatmap: Attention Rollout (Abnar & Zuidema, 2020) — 12 layers × 12 heads")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — TECHNICAL APPROACH (Architecture)
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 5: Technical Approach — Architecture Flow")
s5 = slides[4]

# Change heading
tb12_s5 = find_textbox(s5, 'TextBox 12')
if tb12_s5:
    clear_text(tb12_s5)
    add_line(tb12_s5.text_frame, "SYSTEM ARCHITECTURE", size=HEADING_SIZE, bold=True, first=True)

arch_box = add_textbox(s5, 1790710, 2200000, 14706581, 7000000)
tf = arch_box.text_frame
tf.word_wrap = True

add_heading_line(tf, "End-to-End Pipeline:", first=True)
add_line(tf, "", size=BULLET_SIZE)
add_line(tf, "  Image Upload (File / Camera)  →  Preprocessing (Denoise + Resize)", size=DETAIL_SIZE)
add_line(tf, "      ↓", size=DETAIL_SIZE)
add_line(tf, "  Image Type Detection (H&E Stain / Clinical)  →  ViT Inference (7-class)", size=DETAIL_SIZE)
add_line(tf, "      ↓                                              ↓", size=DETAIL_SIZE)
add_line(tf, "  Attention Rollout Heatmap (14×14 grid)    Gemini Agentic Report", size=DETAIL_SIZE)
add_line(tf, "      ↓                                              ↓", size=DETAIL_SIZE)
add_line(tf, "  Results Dashboard  ←  MongoDB Record  ←  JSON API Response", size=DETAIL_SIZE)
add_line(tf, "", size=BULLET_SIZE)

add_heading_line(tf, "API Endpoints:")
add_bullet(tf, "POST /predict — Main inference (ViT + Heatmap + Gemini Report)")
add_bullet(tf, "POST /api/simulate_progression — Stable Diffusion LoRA melanoma simulation")
add_bullet(tf, "POST /api/chat — Context-aware AI chatbot for follow-up questions")
add_bullet(tf, "GET /api/history/<patient_id> — Patient longitudinal risk tracking")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — TECHNICAL APPROACH (Unique Features)
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 6: Technical Approach — Unique Features")
s6 = slides[5]

tb12_s6 = find_textbox(s6, 'TextBox 12')
if tb12_s6:
    clear_text(tb12_s6)
    add_line(tb12_s6.text_frame, "KEY DIFFERENTIATORS", size=HEADING_SIZE, bold=True, first=True)

feat_box = add_textbox(s6, 1790710, 2200000, 14706581, 7000000)
tf = feat_box.text_frame
tf.word_wrap = True

add_heading_line(tf, "1. ViT Attention Rollout Heatmaps (Explainable AI)", first=True)
add_bullet(tf, "Not Grad-CAM — uses real transformer attention weights across all 12 layers")
add_bullet(tf, "CLS→patch attention extracted, reshaped to 14×14, overlaid with JET colormap")
add_bullet(tf, "Doctors see exactly WHERE the model focuses — builds clinical trust")
add_line(tf, "", size=BULLET_SIZE)

add_heading_line(tf, "2. Agentic AI Medical Reporting (Gemini 2.5 Flash)")
add_bullet(tf, "Dynamic, context-aware reports — not static templates")
add_bullet(tf, "Outputs: medical advice, preventive routine, lifestyle suggestions")
add_bullet(tf, "Conversational follow-up chat with full diagnostic context retention")
add_line(tf, "", size=BULLET_SIZE)

add_heading_line(tf, "3. Melanoma Progression Simulation (Stable Diffusion + LoRA)")
add_bullet(tf, "Custom-trained LoRA adapter on melanoma pathology images")
add_bullet(tf, "Img2Img pipeline: takes benign lesion → simulates potential melanoma development")
add_bullet(tf, "Helps patients visualize risk — powerful motivator for early clinical intervention")
add_line(tf, "", size=BULLET_SIZE)

add_heading_line(tf, "4. Automatic H&E Stain Detection")
add_bullet(tf, "HSV color analysis + texture entropy classifies pathology vs clinical images")
add_bullet(tf, "System adapts recommendations based on image source type")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — RESULTS / EFFICIENCY
# ══════════════════════════════════════════════════════════════════════════════
print("📝 Slide 7: Results & Impact")
s7 = slides[6]

tb13_s7 = find_textbox(s7, 'TextBox 13')
if tb13_s7:
    clear_text(tb13_s7)
    tf = tb13_s7.text_frame
    tf.word_wrap = True

    add_heading_line(tf, "Performance Metrics", first=True)
    add_bullet(tf, "Model Accuracy: 98.52% on HAM10000 test set (7-class)")
    add_bullet(tf, "Inference Speed: <1 second on GPU (RTX 3050), ~4s on CPU")
    add_bullet(tf, "Heatmap Generation: Real-time attention rollout in <500ms")
    add_bullet(tf, "LLM Response: Agentic report generated in ~2 seconds via Gemini Flash")
    add_line(tf, "", size=BULLET_SIZE)

    add_heading_line(tf, "Projected Impact")
    add_bullet(tf, "Early detection of melanoma can improve 5-year survival from 30% to 99%")
    add_bullet(tf, "Reduces diagnostic wait time from days/weeks to seconds")
    add_bullet(tf, "Explainable AI builds doctor trust — accelerates clinical adoption")
    add_line(tf, "", size=BULLET_SIZE)

    add_heading_line(tf, "Scalability Potential")
    add_bullet(tf, "Cloud deployment (AWS/GCP) — serve millions of predictions per day")
    add_bullet(tf, "Multi-disease expansion: retinopathy, oral cancer, lung pathology")
    add_bullet(tf, "Mobile app integration for telemedicine in rural areas")
    add_line(tf, "", size=BULLET_SIZE)

    add_heading_line(tf, "Practical Implementation Value")
    add_bullet(tf, "Hospital integration as second-opinion screening tool")
    add_bullet(tf, "Patient self-screening kiosk for early warning")
    add_bullet(tf, "Medical education platform for training dermatology students")

# ══════════════════════════════════════════════════════════════════════════════
# DELETE SLIDE 8 (Instructions)
# ══════════════════════════════════════════════════════════════════════════════
print("🗑️  Removing instruction slide (slide 8)")
rId = prs.slides._sldIdLst[-1].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
if rId is None:
    # Try alternate attribute
    rId_elem = prs.slides._sldIdLst[-1]
    for attr_name in rId_elem.attrib:
        if 'id' in attr_name.lower() and attr_name != 'id':
            rId = rId_elem.get(attr_name)
            break

# Remove last slide element from the slide list
last_sldId = prs.slides._sldIdLst[-1]
prs.slides._sldIdLst.remove(last_sldId)

# ── Save ─────────────────────────────────────────────────────────────────────
import io
buf = io.BytesIO()
prs.save(buf)
buf.seek(0)
with open(OUTPUT, 'wb') as f:
    f.write(buf.read())

print(f"\n{'='*55}")
print(f"✅ Presentation saved to:\n   {OUTPUT}")
print(f"{'='*55}")
print(f"\n⚠️  Don't forget to replace '[YOUR TEAM NAME]' on Slide 1!")
print(f"💡 Save as PDF: Open in PowerPoint → File → Save As → PDF")
