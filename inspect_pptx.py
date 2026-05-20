"""Inspect the PPT template: list every slide, shape, text, position, and font."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import os

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
    "PPT TEMPLATE - National Level Project Competition - VIT, Pune (1).pptx")

prs = Presentation(TEMPLATE)
print(f"Slide width : {prs.slide_width}  ({prs.slide_width / 914400:.2f} in)")
print(f"Slide height: {prs.slide_height} ({prs.slide_height / 914400:.2f} in)")
print(f"Total slides: {len(prs.slides)}\n")

for si, slide in enumerate(prs.slides, 1):
    layout_name = slide.slide_layout.name if slide.slide_layout else "None"
    print(f"{'='*70}")
    print(f"SLIDE {si}  (layout: {layout_name})")
    print(f"{'='*70}")
    for shape in slide.shapes:
        print(f"\n  Shape: {shape.shape_type} | Name: '{shape.name}'")
        print(f"    Pos : left={shape.left}, top={shape.top}, w={shape.width}, h={shape.height}")
        if shape.has_text_frame:
            for pi, para in enumerate(shape.text_frame.paragraphs):
                txt = para.text.strip()
                if txt:
                    runs_info = []
                    for r in para.runs:
                        fi = f"sz={r.font.size}" if r.font.size else ""
                        fi += f" bold={r.font.bold}" if r.font.bold else ""
                        fi += f" color={r.font.color.rgb}" if r.font.color and r.font.color.rgb else ""
                        fi += f" name={r.font.name}" if r.font.name else ""
                        runs_info.append(fi)
                    print(f"    P{pi}: '{txt[:80]}' [{'; '.join(runs_info)}]")
    print()
