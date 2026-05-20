"""
export_word.py - Export Digital Pathology architecture & workflows to Word
using structured text, tables, and numbered steps (no images needed).

Requirements: pip install python-docx
Output: Digital_Pathology_Report.docx  (in project folder)
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io

OUTPUT_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Digital_Pathology_Report.docx")

# ── Style helpers ──────────────────────────────────────────────────────────────

def set_col_width(cell, width_inches):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def shade_cell(cell, hex_color="1F3864"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        shade_cell(cell, "1F3864")
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        if col_widths:
            set_col_width(cell, col_widths[i])
    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            if (ri % 2) == 0:
                shade_cell(cell, "EBF3FB")
            if col_widths:
                set_col_width(cell, col_widths[ci])
    doc.add_paragraph()

def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

def add_h3(doc, text):
    doc.add_heading(text, level=3)

def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)

def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(step).font.size = Pt(10)

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item).font.size = Pt(10)

def page_break(doc):
    doc.add_page_break()

# ── Build document ─────────────────────────────────────────────────────────────

def build():
    doc = Document()

    # Page margins
    sec = doc.sections[0]
    sec.left_margin  = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin   = Inches(1.0)
    sec.bottom_margin = Inches(1.0)

    # ── Cover ──────────────────────────────────────────────────────────────────
    t = doc.add_heading("Digital Pathology AI System", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Architecture & Workflow Documentation")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14)

    doc.add_paragraph()
    add_table(doc,
        ["Property", "Value"],
        [
            ["Project",      "Digital Pathology Image Analysis"],
            ["Model",        "google/vit-base-patch16-224 (fine-tuned, HAM10000)"],
            ["Accuracy",     "98.52% on 7-class skin cancer classification"],
            ["Framework",    "Flask + PyTorch + Stable Diffusion + Gemini 2.5 Flash"],
            ["Database",     "MongoDB (pathai_db)"],
            ["Hardware",     "CUDA (RTX 3050 6GB) / CPU fallback"],
        ],
        col_widths=[2.0, 4.5]
    )
    page_break(doc)

    # ══════════════════════════════════════════════════════════════════════════
    # PART 1 — ARCHITECTURE
    # ══════════════════════════════════════════════════════════════════════════
    add_h1(doc, "Part 1: System Architecture")
    add_para(doc, "A full-stack AI web application for skin cancer analysis using a fine-tuned Vision Transformer, a Stable Diffusion generative engine, and a Gemini LLM agentic layer.")

    # 1.1 High-Level Layers
    add_h2(doc, "1.1  High-Level System Layers")
    add_table(doc,
        ["Layer", "Technology", "Description"],
        [
            ["Frontend",          "HTML + Vanilla JS + CSS3",        "Single-page app: upload, results, heatmap, chat, history"],
            ["Backend",           "Flask (Python)",                   "REST API server handling routing and orchestration"],
            ["AI Inference",      "PyTorch + HuggingFace Transformers","ViT classifier + attention rollout heatmap engine"],
            ["Generative Engine", "Stable Diffusion v1.5 + LoRA",    "Melanoma progression simulation via img2img pipeline"],
            ["Agentic AI",        "Google Gemini 2.5 Flash",          "Medical report generation and conversational chat"],
            ["Database",          "MongoDB",                          "Patient record persistence and history tracking"],
            ["Storage",           "File System (static/uploads/)",    "Uploaded and generated image files"],
        ],
        col_widths=[1.5, 2.2, 3.0]
    )

    # 1.2 Frontend Files
    add_h2(doc, "1.2  Frontend Components  (templates/ + static/)")
    add_table(doc,
        ["File", "Role"],
        [
            ["templates/index.html",   "Single-page app: file upload, camera capture, results dashboard, heatmap viewer, AI chat, patient history chart"],
            ["static/script.js",       "All client-side logic: fetch API calls, Chart.js probability bar chart, base64 image handling, camera capture"],
            ["static/style.css",       "Dark-mode glassmorphism UI, animated gradients, micro-animations"],
            ["static/background.mp4",  "Animated video background rendered behind the UI"],
            ["static/uploads/",        "Server-side storage for uploaded and generated images"],
        ],
        col_widths=[2.2, 4.5]
    )

    # 1.3 Flask API Routes
    add_h2(doc, "1.3  Flask API Routes  (app.py)")
    add_table(doc,
        ["Route", "Method", "Description"],
        [
            ["/",                           "GET",  "Serves index.html (main application page)"],
            ["/predict",                    "POST", "Main inference: accepts file or base64 image, returns prediction + heatmap + report"],
            ["/api/simulate_progression",   "POST", "Stable Diffusion LoRA img2img: simulates melanoma progression on a benign lesion"],
            ["/api/chat",                   "POST", "Gemini-powered conversational Q&A with patient diagnostic context"],
            ["/api/history/<patient_id>",   "GET",  "Fetches chronological analysis records for a patient from MongoDB"],
        ],
        col_widths=[2.5, 0.8, 3.5]
    )

    # 1.4 ViT Model
    add_h2(doc, "1.4  ViT Classifier  (models.py + ./saved_model)")
    add_table(doc,
        ["Property", "Value"],
        [
            ["Base Model",     "google/vit-base-patch16-224"],
            ["Training Data",  "HAM10000 (marmal88/skin_cancer @ HuggingFace)"],
            ["Accuracy",       "98.52%"],
            ["Input Size",     "224 x 224 RGB, normalized: mean=0.5, std=0.5"],
            ["Output",         "7-class softmax probabilities"],
            ["Inference",      "CUDA (RTX 3050) with CPU fallback"],
            ["Saved Files",    "config.json, model.safetensors (~327 MB), preprocessor_config.json"],
        ],
        col_widths=[2.0, 4.5]
    )

    add_h3(doc, "7 HAM10000 Classes and Risk Levels")
    add_table(doc,
        ["Class ID", "Display Name", "Category", "Risk Level"],
        [
            ["actinic_keratoses",           "Actinic Keratoses",   "Pre-Malignant", "Medium"],
            ["basal_cell_carcinoma",         "Basal Cell Carcinoma","Malignant",     "High"],
            ["benign_keratosis-like_lesions","Benign Keratosis",    "Benign",        "Low"],
            ["dermatofibroma",               "Dermatofibroma",      "Benign",        "Low"],
            ["melanocytic_Nevi",             "Melanocytic Nevi",    "Benign",        "Low"],
            ["melanoma",                     "Melanoma",            "Malignant",     "High"],
            ["vascular_lesions",             "Vascular Lesions",    "Benign",        "Low"],
        ],
        col_widths=[2.2, 1.8, 1.4, 1.2]
    )

    # 1.5 Heatmap
    add_h2(doc, "1.5  Heatmap Engine  (heatmaps.py)")
    add_para(doc, "3-tier fallback strategy to generate a saliency map regardless of model type:")
    add_table(doc,
        ["Priority", "Method", "Condition", "Algorithm"],
        [
            ["1 (Best)",  "ViT Attention Rollout", "Model is _ViTWrapper",    "Abnar & Zuidema 2020: accumulate 12-layer attention, extract CLS→patch 14x14 grid"],
            ["2",         "Grad-CAM",              "Model has Conv2D layers", "Gradients of top class w.r.t. last Conv2D feature map, global avg pooled"],
            ["3 (Fallback)", "Simulated Edge-Energy","All other cases",        "Laplacian edge detection + centred Gaussian blend, normalized to [0,1]"],
        ],
        col_widths=[0.9, 1.7, 1.7, 2.5]
    )
    add_para(doc, "All methods output a [0,1] saliency map which is resized to the original image dimensions, colorized with COLORMAP_JET (blue=low, red=high), and blended 50/50 with the original image.")

    # 1.6 Generative Engine
    add_h2(doc, "1.6  Generative Engine  (Stable Diffusion + LoRA)")
    add_table(doc,
        ["Property", "Value"],
        [
            ["Base Model",     "Stable Diffusion v1.5 (v1-5-pruned-emaonly.safetensors)"],
            ["LoRA Weights",   "PathAI_LoRA/model/pathai_melanoma_v1.safetensors"],
            ["Pipeline",       "StableDiffusionImg2ImgPipeline (diffusers)"],
            ["Input Size",     "512 x 512 RGB"],
            ["Strength",       "0.65 (preserves skin structure, adds melanoma texture)"],
            ["Guidance Scale", "7.5"],
            ["VRAM Strategy",  "enable_model_cpu_offload() for RTX 3050 6GB"],
            ["Prompt",         "macro clinical photography of a p4th4i melanoma on skin, highly detailed, sharp focus, medical imaging"],
        ],
        col_widths=[2.0, 4.5]
    )

    # 1.7 Agentic AI
    add_h2(doc, "1.7  Agentic AI Layer  (utils.py + Gemini)")
    add_table(doc,
        ["Function", "Model", "Input", "Output"],
        [
            ["generate_agentic_report()", "Gemini 2.5 Flash", "prediction, confidence, risk_level, image_type", "JSON: medical_advice, preventive_routine, lifestyle_suggestions, maps_link, disclaimer"],
            ["chat_with_agent()",         "Gemini 2.5 Flash", "user_message, context_summary, chat_history",    "Plain-text conversational response with safety constraints"],
        ],
        col_widths=[2.0, 1.5, 2.0, 2.5]
    )

    # 1.8 Database
    add_h2(doc, "1.8  Database Schema  (MongoDB)")
    add_table(doc,
        ["Collection", "Field", "Type", "Description"],
        [
            ["pathai_db.patients", "patient_id",  "String",   "Patient identifier (default: 'Anonymous')"],
            ["pathai_db.patients", "date",         "DateTime", "UTC timestamp of analysis"],
            ["pathai_db.patients", "image_path",   "String",   "Server-side path of uploaded image"],
            ["pathai_db.patients", "prediction",   "String",   "Predicted class ID (e.g. melanoma)"],
            ["pathai_db.patients", "confidence",   "Float",    "Model confidence percentage (0-100)"],
            ["pathai_db.patients", "risk_level",   "String",   "Low / Medium / High"],
        ],
        col_widths=[1.8, 1.2, 1.0, 3.2]
    )

    # 1.9 Training Pipeline
    add_h2(doc, "1.9  Training Pipeline  (Offline Scripts)")
    add_table(doc,
        ["Script", "Purpose", "Output"],
        [
            ["train.py",       "Fine-tune google/vit-base-patch16-224 on HAM10000 7 classes",     "./saved_model/ (config + weights + processor)"],
            ["finetune.py",    "LoRA fine-tuning of Stable Diffusion on melanoma images",          "PathAI_LoRA/model/pathai_melanoma_v1.safetensors"],
            ["prep_dataset.py","Curate and format melanoma images from raw_melanoma_selection/",   "Processed dataset for LoRA training"],
            ["ImportModel.py", "Model import and conversion utilities",                            "Imported model weights"],
            ["test_lora.py",   "Test the LoRA fine-tuned generative model",                       "Test output images"],
            ["test_rollout.py","Validate ViT attention rollout heatmap output",                    "Heatmap PNGs for visual inspection"],
            ["test_api.py",    "End-to-end API test against the running Flask server",             "API response validation"],
        ],
        col_widths=[1.5, 3.0, 2.2]
    )

    page_break(doc)

    # ══════════════════════════════════════════════════════════════════════════
    # PART 2 — WORKFLOWS
    # ══════════════════════════════════════════════════════════════════════════
    add_h1(doc, "Part 2: Workflow Descriptions")

    # WF1
    add_h2(doc, "Workflow 1: Image Analysis Pipeline  (POST /predict)")
    add_para(doc, "Core end-to-end flow from image upload to full clinical report:")
    add_steps(doc, [
        "User selects input method: file upload (multipart) or camera capture (base64 JSON).",
        "Optionally enters a Patient ID (defaults to 'Anonymous').",
        "Frontend sends POST /predict with the image data.",
        "Flask validates: file extension in {png, jpg, jpeg, bmp, tiff, webp}, size <= 16 MB, dimensions >= 32x32 px. Returns 400 on failure.",
        "Valid image saved to static/uploads/<uuid>_filename.",
        "preprocess_image(): BGR→RGB conversion, fastNlMeansDenoising, LANCZOS resize to 224x224, normalize to [0,1], add batch dimension → shape (1,224,224,3).",
        "detect_image_type(): HSV hue analysis for H&E stain (purple/pink ratio), saturation std deviation, grayscale histogram entropy → returns 'pathology' or 'external'.",
        "_ViTWrapper.predict(): loads PIL image from original file, LANCZOS resize, ViT normalization (mean=0.5, std=0.5), HWC→CHW tensor, forward pass → logits (1x7), softmax → probabilities.",
        "Returns: top_class, display_name, confidence, risk_level (Low/Medium/High), category (Benign/Pre-Malignant/Malignant), all_probs dict.",
        "generate_heatmap(): runs Attention Rollout (preferred) → Grad-CAM → Simulated edge-energy fallback → returns base64 JPEG overlay.",
        "generate_agentic_report(): calls Gemini 2.5 Flash with prediction context → returns JSON with medical_advice, preventive_routine, lifestyle_suggestions, maps_link.",
        "MongoDB insert: saves patient_id, date, image_path, prediction, confidence, risk_level to pathai_db.patients.",
        "JSON response returned: prediction, display_name, confidence, risk_level, all_probs, heatmap (base64), original_image (base64), recommendation.",
        "Frontend renders: risk badge, confidence %, Chart.js probability bar chart, heatmap overlay, medical advice cards, preventive steps list, Google Maps link.",
    ])

    # WF2
    add_h2(doc, "Workflow 2: Progression Simulation  (POST /api/simulate_progression)")
    add_para(doc, "Uses custom-trained LoRA over Stable Diffusion to simulate melanoma development:")
    add_steps(doc, [
        "User uploads a benign skin lesion image and clicks 'Simulate Progression'.",
        "Flask checks gen_pipe is loaded; returns 503 if the generative engine is offline.",
        "File validation: extension check and non-empty filename.",
        "PIL opens image, converts to RGB, resizes to 512x512.",
        "Clinical prompt constructed: 'macro clinical photography of a p4th4i melanoma on skin, highly detailed, sharp focus, medical imaging'.",
        "StableDiffusionImg2ImgPipeline runs with LoRA weights loaded, strength=0.65, guidance_scale=7.5, enable_model_cpu_offload() active for VRAM management.",
        "Generated 512x512 synthetic melanoma image saved to static/uploads/simulated_<uuid>.jpg.",
        "Image encoded to base64 and returned in JSON response.",
        "Frontend displays the synthetic progression image alongside the original.",
    ])

    # WF3
    add_h2(doc, "Workflow 3: AI Chat  (POST /api/chat)")
    add_para(doc, "Context-aware conversational assistant pre-loaded with patient diagnostic context:")
    add_steps(doc, [
        "User types a question in the chat panel after a diagnosis has been performed.",
        "script.js sends POST /api/chat with: message (user question), context (diagnosis summary), history (previous chat turns).",
        "Flask validates message is not empty; returns 400 otherwise.",
        "System prompt built: expert dermatological AI role + patient context + conversation history + current question + safety constraints (no definitive diagnosis, always recommend doctor).",
        "Gemini 2.5 Flash called via client.models.generate_content().",
        "On success: response.text.strip() returned as plain-text reply.",
        "On failure: fallback message returned: 'Having trouble connecting. Please consult your doctor.'",
        "JSON response with reply field sent to frontend.",
        "Frontend appends the AI reply as a new message bubble in the chat window.",
    ])

    # WF4
    add_h2(doc, "Workflow 4: Patient History  (GET /api/history/<patient_id>)")
    add_para(doc, "Retrieves and visualizes chronological analysis records for a patient:")
    add_steps(doc, [
        "User enters a Patient ID in the history panel.",
        "script.js sends GET /api/history/<patient_id>.",
        "MongoDB query: db.patients.find(patient_id).sort(date, ascending=1).",
        "Each record formatted: date → 'Apr 20, 2025 - 14:32', prediction label, confidence %, risk_level.",
        "JSON response returned with history array.",
        "script.js renders a Chart.js line graph: X-axis = dates, Y-axis = confidence %, points color-coded by risk level.",
        "Clinician can visually track risk progression over time.",
    ])

    # WF5
    add_h2(doc, "Workflow 5: Model Training Pipeline  (Offline)")
    add_h3(doc, "5a. ViT Fine-Tuning  (train.py)")
    add_steps(doc, [
        "Load HAM10000 dataset via load_dataset('marmal88/skin_cancer') from HuggingFace Hub.",
        "Define 7 class labels: akiec, bcc, bkl, df, mel, nv, vasc with label2id/id2label mappings.",
        "Load ViTForImageClassification from google/vit-base-patch16-224 with num_labels=7, ignore_mismatched_sizes=True (replaces the 1000-class head).",
        "Fine-tune on HAM10000 → achieves 98.52% accuracy.",
        "Save via save_pretrained('./saved_model') → produces config.json, model.safetensors, preprocessor_config.json.",
    ])
    add_h3(doc, "5b. LoRA Fine-Tuning  (finetune.py)")
    add_steps(doc, [
        "prep_dataset.py curates melanoma images from raw_melanoma_selection/ folder.",
        "LoRA adapters trained over Stable Diffusion v1.5 base model on melanoma image dataset.",
        "Weights saved to PathAI_LoRA/model/pathai_melanoma_v1.safetensors.",
        "Loaded at runtime via gen_pipe.load_lora_weights().",
    ])

    # WF6
    add_h2(doc, "Workflow 6: Heatmap Generation Decision Tree  (heatmaps.py)")
    add_steps(doc, [
        "generate_heatmap(model, preprocessed_array, original_path) called after inference.",
        "cv2.imread(original_path) — if unreadable, return blank grey placeholder.",
        "Check if model has .model, .processor, .device attributes (i.e., is a _ViTWrapper).",
        "  YES → run _attention_rollout(): build cached eager-mode ViT, forward pass with output_attentions=True, stack 12 layers (L,H,S,S), average over heads, add identity matrix (residual skip), row-normalize, matrix-multiply all layers (rollout), extract CLS→patch row [0,1:] → reshape to 14x14, normalize to [0,1]. If exception → fall through.",
        "  NO → check if model has Conv2D layers and TensorFlow is available.",
        "  YES → run _grad_cam(): find last Conv2D, build GradientTape model, compute gradients of top class, global avg pool → weighted feature map sum → ReLU → normalize. If exception → fall through.",
        "  NO (or fallback) → run _simulated_cam(): Laplacian edge detection, Gaussian blur (sigma=20), centred Gaussian blend (0.6 x edge + 0.4 x centre), normalize to [0,1].",
        "Resize saliency map to original image dimensions using INTER_CUBIC + GaussianBlur.",
        "Apply COLORMAP_JET (blue=low attention, red=high attention).",
        "Blend 50% original image + 50% colorized heatmap via addWeighted.",
        "Draw colour scale legend (Low | Attention | High) in bottom-right corner.",
        "Encode as JPEG (quality=90) and return as base64 data URL.",
    ])

    # WF7 — connections table
    add_h2(doc, "Workflow 7: Full System Interaction Map")
    add_para(doc, "How all user actions connect to backend routes and AI components:")
    add_table(doc,
        ["User Action", "Route", "AI Components Used", "Storage"],
        [
            ["Upload image for analysis",   "POST /predict",                    "ViT Classifier, Attention Rollout Heatmap, Gemini 2.5 Flash", "MongoDB + File System"],
            ["Simulate progression",        "POST /api/simulate_progression",   "Stable Diffusion v1.5 + PathAI LoRA",                        "File System"],
            ["Ask AI question",             "POST /api/chat",                   "Gemini 2.5 Flash (with context)",                            "None"],
            ["View patient history",        "GET /api/history/<patient_id>",    "None (data retrieval only)",                                 "MongoDB"],
            ["Load main page",              "GET /",                            "None",                                                       "None"],
        ],
        col_widths=[1.8, 2.0, 2.5, 1.3]
    )

    # ── Save ──────────────────────────────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    with open(OUTPUT_FILE, 'wb') as f:
        f.write(buf.read())

    print(f"\n{'='*55}")
    print(f"✅ Word document saved to:\n   {OUTPUT_FILE}")
    print(f"{'='*55}")

if __name__ == '__main__':
    try:
        from docx import Document
    except ImportError:
        print("❌ Run: pip install python-docx")
        exit(1)
    build()
